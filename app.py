"""AI Resume ATS Checker - Streamlit + Google Gemini Flash."""

import io
import json
import os
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

DEFAULT_MODEL = "gemini-3.5-flash"
MAX_CHARS = 20000  # keep prompts small and fast

PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyze the resume below{jd_clause} and respond with ONLY a JSON object (no markdown, no commentary)
using exactly this schema:

{{
  "ats_score": <integer 0-100>,
  "breakdown": {{
    "formatting": <integer 0-100>,
    "keywords": <integer 0-100>,
    "experience_impact": <integer 0-100>,
    "skills": <integer 0-100>,
    "readability": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<short point>", ...],
  "improvements": [
    {{"issue": "<what is wrong>", "fix": "<specific action to take>", "priority": "high|medium|low"}}
  ],
  "missing_keywords": ["<keyword>", ...],
  "rewrite_examples": [
    {{"before": "<weak line from the resume>", "after": "<stronger rewritten line>"}}
  ]
}}

Scoring guidance: be honest and calibrated; most resumes score 50-80. Penalise missing contact info,
missing sections, vague bullets without metrics, keyword gaps, tables/columns that ATS may misread.
Give 3-6 strengths, 5-8 improvements, up to 10 missing keywords, and 2-4 rewrite examples that use
ONLY facts present in the resume (never invent numbers or employers).
{jd_block}
RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ---------- File reading ----------
def extract_text(uploaded_file) -> str:
    """Return plain text from a PDF, DOCX or TXT upload."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(parts)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")

    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------- AI call ----------
def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def parse_response(raw: str) -> dict:
    """Parse and sanitise the model's JSON reply."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        data = json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if not match:
            raise ValueError("The AI returned an unreadable response. Please try again.")
        data = json.loads(match.group(0))

    breakdown = data.get("breakdown") or {}
    keys = ["formatting", "keywords", "experience_impact", "skills", "readability"]
    return {
        "ats_score": _clamp(data.get("ats_score")),
        "breakdown": {k: _clamp(breakdown.get(k)) for k in keys},
        "summary": str(data.get("summary", "")),
        "strengths": [str(s) for s in data.get("strengths", []) if s],
        "improvements": [i for i in data.get("improvements", []) if isinstance(i, dict)],
        "missing_keywords": [str(k) for k in data.get("missing_keywords", []) if k],
        "rewrite_examples": [r for r in data.get("rewrite_examples", []) if isinstance(r, dict)],
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    jd = (job_description or "").strip()
    prompt = PROMPT.format(
        jd_clause=" against the job description" if jd else "",
        jd_block=f'\nJOB DESCRIPTION:\n"""\n{jd[:MAX_CHARS]}\n"""\n' if jd else "",
        resume=resume_text[:MAX_CHARS],
    )
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return parse_response(response.text)


def friendly_error(exc: Exception) -> str:
    msg = str(exc)
    low = msg.lower()
    if "404" in msg or "not found" in low or "no longer available" in low:
        return ("The selected Gemini model isn't available to your key. Change the model name in "
                "the sidebar (check Google AI Studio for the current Flash model). Details: " + msg)
    if "api key" in low or "401" in msg or "403" in msg or "permission" in low:
        return "Your Gemini API key looks invalid or lacks permission. Details: " + msg
    if "429" in msg or "quota" in low or "rate" in low:
        return "Rate limit or quota reached. Wait a minute and try again. Details: " + msg
    return "Something went wrong while analysing the resume: " + msg


# ---------- UI ----------
def score_label(score: int):
    if score >= 80:
        return "Excellent", "green"
    if score >= 65:
        return "Good", "orange"
    return "Needs work", "red"


def get_api_key(sidebar_value: str) -> str:
    if sidebar_value:
        return sidebar_value
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass
    return os.environ.get("GEMINI_API_KEY", "")


def render_results(result: dict):
    score = result["ats_score"]
    label, color = score_label(score)

    st.subheader("Your ATS score")
    c1, c2 = st.columns([1, 2])
    with c1:
        st.metric("ATS Score", f"{score}/100")
        st.markdown(f":{color}[**{label}**]")
    with c2:
        st.progress(score / 100)
        st.write(result["summary"])

    st.subheader("Score breakdown")
    cols = st.columns(5)
    names = {
        "formatting": "Formatting",
        "keywords": "Keywords",
        "experience_impact": "Impact",
        "skills": "Skills",
        "readability": "Readability",
    }
    for col, (key, title) in zip(cols, names.items()):
        col.metric(title, result["breakdown"][key])

    left, right = st.columns(2)
    with left:
        st.subheader("Strengths")
        for s in result["strengths"] or ["No strengths listed."]:
            st.write(f"- {s}")
    with right:
        st.subheader("Missing keywords")
        if result["missing_keywords"]:
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))
        else:
            st.write("None found.")

    st.subheader("Suggested improvements")
    order = {"high": 0, "medium": 1, "low": 2}
    items = sorted(result["improvements"], key=lambda i: order.get(str(i.get("priority", "")).lower(), 3))
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for item in items:
        pr = str(item.get("priority", "")).lower()
        with st.expander(f"{icons.get(pr, '⚪')} {item.get('issue', 'Improvement')}"):
            st.write(item.get("fix", ""))

    if result["rewrite_examples"]:
        st.subheader("Rewrite examples")
        for ex in result["rewrite_examples"]:
            st.markdown(f"**Before:** {ex.get('before', '')}")
            st.markdown(f"**After:** {ex.get('after', '')}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main():
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and concrete ways to improve it.")

    with st.sidebar:
        st.header("Settings")
        key_input = st.text_input("Gemini API key", type="password",
                                  help="Leave empty if the key is set in Streamlit secrets.")
        model = st.text_input("Model", value=os.environ.get("GEMINI_MODEL", DEFAULT_MODEL))
        st.markdown("Get a free key at [Google AI Studio](https://aistudio.google.com/apikey).")
        st.info("Your resume is sent to Google's Gemini API for analysis and isn't stored by this app.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    jd = st.text_area("Job description (optional, for a tailored score)", height=150)

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        api_key = get_api_key(key_input)
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar or in Streamlit secrets.")
            return
        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
            if len(text) < 100:
                st.error("Couldn't read enough text. If this is a scanned PDF, upload a text-based "
                         "PDF or DOCX (ATS systems can't read scanned images either).")
                return
            with st.spinner("Analysing with Gemini..."):
                result = analyze_resume(text, jd, api_key, model.strip() or DEFAULT_MODEL)
            st.session_state["result"] = result
        except Exception as exc:  # noqa: BLE001
            st.error(friendly_error(exc))
            return

    if "result" in st.session_state:
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
